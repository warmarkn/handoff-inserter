"""Generate fixtures/synthetic-handoff.docx: fake patients only, no PHI.

Covers the spec's edge cases:
  - beds 31-72 with 44 deliberately missing
  - stray "40 mg ..." line inside bed 35 (breaks ascending order)
  - stray "38 hr NPO" line inside bed 37 (same LIS length; formatting breaks the tie)
  - trailing blank paragraphs after bed 33's body
  - bulleted body text in bed 41
  - bed 47 with no body paragraph
  - beds 50-55 inside a table (one row per bed)
  - bed 60 inside a VML text box
  - bed 72 is the last paragraph in the document
"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml

OUT = os.path.join(os.path.dirname(__file__), "..", "fixtures", "synthetic-handoff.docx")
NAMES = ["Doe", "Roe", "Poe", "Moe", "Loe", "Koe", "Zoe", "Hoe"]


def bed_line(n):
    return f"{n} {NAMES[n % len(NAMES)]}, Test  {40 + n % 40}{'MF'[n % 2]}  synthetic dx {n}"


def add_bed(container, n, body=True):
    p = container.add_paragraph()
    r = p.add_run(bed_line(n))
    r.bold = True
    r.font.size = Pt(11)
    if body:
        b = container.add_paragraph()
        br = b.add_run(f"HPI: fake history for bed {n}.")
        br.font.size = Pt(10)
        b.paragraph_format.left_indent = Pt(18)
    return p


def body_line(container, text, style=None):
    b = container.add_paragraph(style=style)
    r = b.add_run(text)
    r.font.size = Pt(10)
    b.paragraph_format.left_indent = Pt(18)
    return b


def textbox_paragraph(doc, n):
    """A paragraph whose run holds a VML text box containing a bed block."""
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    V = 'xmlns:v="urn:schemas-microsoft-com:vml"'
    xml = f"""
<w:p {W} {V}>
  <w:r>
    <w:pict>
      <v:shape style="width:300pt;height:60pt">
        <v:textbox>
          <w:txbxContent>
            <w:p><w:r><w:rPr><w:b/><w:sz w:val="22"/></w:rPr><w:t xml:space="preserve">{bed_line(n)}</w:t></w:r></w:p>
            <w:p><w:pPr><w:ind w:left="360"/></w:pPr><w:r><w:rPr><w:sz w:val="20"/></w:rPr><w:t xml:space="preserve">HPI: fake history for bed {n} (text box).</w:t></w:r></w:p>
          </w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>
  </w:r>
</w:p>"""
    doc.element.body.sectPr.addprevious(parse_xml(xml))


def main():
    doc = Document()
    doc.add_heading("ICU Handoff (SYNTHETIC TEST DATA)", level=1)

    for n in range(31, 50):
        if n == 44:
            continue
        add_bed(doc, n, body=(n != 47))
        if n == 33:
            doc.add_paragraph("")
            doc.add_paragraph("")
        if n == 35:
            body_line(doc, "40 mg IV furosemide BID")
        if n == 37:
            body_line(doc, "38 hr NPO, advance diet as tolerated")
        if n == 41:
            body_line(doc, "Plan item one", style="List Bullet")
            body_line(doc, "Plan item two", style="List Bullet")

    table = doc.add_table(rows=0, cols=1)
    for n in range(50, 56):
        cell = table.add_row().cells[0]
        cell.paragraphs[0].add_run(bed_line(n)).bold = True
        body_line(cell, f"HPI: fake history for bed {n} (table).")

    for n in range(56, 72):
        if n == 60:
            textbox_paragraph(doc, n)
            continue
        add_bed(doc, n)

    add_bed(doc, 72, body=False)
    # python-docx puts sectPr last in body; make sure bed 72 is the last paragraph.
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("wrote", os.path.abspath(OUT))


if __name__ == "__main__":
    main()
